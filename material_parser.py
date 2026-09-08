import os
import re
import logging

logger = logging.getLogger(__name__)

def extract_text_from_file(file_path: str) -> dict:
    """
    업로드된 수업자료(PDF, PPTX, TXT)에서 텍스트 및 슬라이드/페이지별 구조를 추출합니다.
    """
    if not os.path.exists(file_path):
        return {"success": False, "error": "파일이 존재하지 않습니다."}

    ext = os.path.splitext(file_path)[1].lower()
    filename = os.path.basename(file_path)

    try:
        if ext == '.pdf':
            return _extract_from_pdf(file_path, filename)
        elif ext in ['.pptx', '.ppt']:
            return _extract_from_pptx(file_path, filename)
        elif ext in ['.txt', '.md']:
            return _extract_from_txt(file_path, filename)
        else:
            return {
                "success": False,
                "error": f"지원되지 않는 파일 형식입니다 ({ext}). PDF 또는 PPTX 파일을 업로드해 주세요."
            }
    except Exception as e:
        logger.error(f"수업자료 파싱 오류 ({filename}): {e}", exc_info=True)
        return {
            "success": False,
            "error": f"파일 텍스트 추출 중 오류가 발생했습니다: {str(e)}"
        }

def _extract_from_pdf(file_path: str, filename: str) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    total_pages = len(reader.pages)
    units = []
    full_text_list = []

    for idx, page in enumerate(reader.pages):
        page_num = idx + 1
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            # 첫 번째 의미 있는 줄을 제목 후보로 사용
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            title = lines[0] if lines else f"Page {page_num}"
            units.append({
                "index": page_num,
                "title": title[:60],
                "text": text
            })
            full_text_list.append(f"[페이지 {page_num}: {title[:40]}]\n{text}")

    full_text = "\n\n".join(full_text_list)
    return {
        "success": True,
        "filename": filename,
        "file_type": "pdf",
        "total_units": total_pages,
        "full_text": full_text,
        "summary_snippet": full_text[:1200] if full_text else "내용 없음",
        "units": units
    }

def _extract_from_pptx(file_path: str, filename: str) -> dict:
    from pptx import Presentation
    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    units = []
    full_text_list = []

    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        slide_texts = []
        slide_title = ""

        # 슬라이드 내 모든 도형에서 텍스트 수집
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)

        # 발표자 노트도 포함 (있다면)
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_texts.append(f"(발표자 노트: {notes})")
        except Exception:
            pass

        if slide_texts:
            slide_title = slide_texts[0]
            joined_text = "\n".join(slide_texts)
            units.append({
                "index": slide_num,
                "title": slide_title[:60],
                "text": joined_text
            })
            full_text_list.append(f"[슬라이드 {slide_num}: {slide_title[:40]}]\n{joined_text}")

    full_text = "\n\n".join(full_text_list)
    return {
        "success": True,
        "filename": filename,
        "file_type": "pptx",
        "total_units": total_slides,
        "full_text": full_text,
        "summary_snippet": full_text[:1200] if full_text else "내용 없음",
        "units": units
    }

def _extract_from_txt(file_path: str, filename: str) -> dict:
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read().strip()

    paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
    units = [{"index": i+1, "title": p[:40], "text": p} for i, p in enumerate(paragraphs)]

    return {
        "success": True,
        "filename": filename,
        "file_type": "txt",
        "total_units": len(paragraphs) or 1,
        "full_text": text,
        "summary_snippet": text[:1200],
        "units": units
    }
